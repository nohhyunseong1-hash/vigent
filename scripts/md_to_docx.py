#!/usr/bin/env python3
"""[문서] md → docx 변환 — python-docx 기반 (v2).

★내력
  v1(2026-09-03 오전)은 python-docx 미설치 상태라 zip+XML 을 손으로 조립했다.
  같은 날 사용자 승인으로 python-docx 를 설치해(requirements-optional.txt)
  라이브러리 기반으로 재작성했다 — 관계(rels)·네임스페이스를 라이브러리가 관리하므로
  손조립보다 워드프로세서 호환성이 안전하다.

★양식 계승
  --template 로 받은 양식 docx 를 **열어 본문만 비우고** 내용을 다시 채운다.
  → 양식 파일의 스타일 정의·기본 글꼴·용지 설정(sectPr)을 그대로 물려받는다.

지원 문법(이 저장소 md 부분집합 — 범위 밖 문법은 일반 문단으로 떨어진다):
  # ~ ### 제목 · GFM 표 · '- ' 목록(2칸 들여쓰기 1단계) · '> ' 인용 ·
  **굵게** · `코드` · [글](링크) · '---' 구분선 무시 · '[공란 …]' 노란 형광펜

사용:
  python scripts/md_to_docx.py --md <입력.md> --template <양식.docx> --out <출력.docx>
"""
from __future__ import annotations

import argparse
import re
import sys
from pathlib import Path

if hasattr(sys.stdout, "reconfigure"):
    sys.stdout.reconfigure(encoding="utf-8")

HL_RE = re.compile(r"(\[공란[^\]]*\])")            # 대표가 채울 칸 — 노란 형광펜
LINK_RE = re.compile(r"\[([^\]]+)\]\((https?://[^)]+)\)")
BOLD_RE = re.compile(r"\*\*(.+?)\*\*")
CODE_RE = re.compile(r"`([^`]+)`")
EAST, WEST = "맑은 고딕", "Malgun Gothic"


def main() -> int:
    import docx
    from docx.enum.text import WD_COLOR_INDEX
    from docx.opc.constants import RELATIONSHIP_TYPE as RT
    from docx.oxml.ns import qn
    from docx.shared import Pt, RGBColor

    ap = argparse.ArgumentParser()
    ap.add_argument("--md", required=True)
    ap.add_argument("--template", required=True, help="★양식 docx — 스타일·용지 설정을 물려받는다")
    ap.add_argument("--out", required=True)
    args = ap.parse_args()
    md_p, tpl, out = Path(args.md), Path(args.template), Path(args.out)
    for p in (md_p, tpl):
        if not p.exists():
            print(f"❌ 없음: {p}")
            return 1

    doc = docx.Document(str(tpl))

    # ── 양식 본문 비우기(용지 설정 sectPr 만 남긴다) ─────────────
    body = doc.element.body
    for child in list(body):
        if not child.tag.endswith("}sectPr"):
            body.remove(child)

    # ── 헬퍼 ────────────────────────────────────────────────────
    def style_run(r, sz, bold=False, color=None, hl=False, code=False):
        r.font.name = "Consolas" if code else WEST
        r.element.rPr.rFonts.set(qn("w:eastAsia"), EAST)
        r.font.size = Pt(sz)
        r.bold = bold
        if color:
            r.font.color.rgb = RGBColor.from_string(color)
        if hl:
            r.font.highlight_color = WD_COLOR_INDEX.YELLOW

    def add_link(par, text, url, sz):
        """python-docx 에는 고수준 하이퍼링크 API 가 없다 — 관계 등록 + oxml 로 만든다."""
        rid = par.part.relate_to(url, RT.HYPERLINK, is_external=True)
        h = docx.oxml.OxmlElement("w:hyperlink")
        h.set(qn("r:id"), rid)
        r = docx.oxml.OxmlElement("w:r")
        rpr = docx.oxml.OxmlElement("w:rPr")
        for tag, attrs in (("w:rFonts", {"w:ascii": WEST, "w:hAnsi": WEST, "w:eastAsia": EAST}),
                           ("w:color", {"w:val": "0563C1"}),
                           ("w:sz", {"w:val": str(sz * 2)}),
                           ("w:u", {"w:val": "single"})):
            e = docx.oxml.OxmlElement(tag)
            for k, v in attrs.items():
                e.set(qn(k), v)
            rpr.append(e)
        t = docx.oxml.OxmlElement("w:t")
        t.set(qn("xml:space"), "preserve")
        t.text = text
        r.append(rpr)
        r.append(t)
        h.append(r)
        par._p.append(h)

    def inline(par, text, sz, bold=False, color=None):
        """**굵게**·`코드`·[링크]·[공란 형광펜] 을 런으로 풀어 넣는다."""
        pos = 0
        for lm in LINK_RE.finditer(text):
            _plain(par, text[pos:lm.start()], sz, bold, color)
            add_link(par, lm.group(1), lm.group(2), sz)
            pos = lm.end()
        _plain(par, text[pos:], sz, bold, color)

    def _plain(par, text, sz, bold, color):
        for i, seg in enumerate(HL_RE.split(text)):
            hl = bool(i % 2)
            p2 = 0
            for bm in BOLD_RE.finditer(seg):
                _code(par, seg[p2:bm.start()], sz, bold, color, hl)
                _code(par, bm.group(1), sz, True, color, hl)
                p2 = bm.end()
            _code(par, seg[p2:], sz, bold, color, hl)

    def _code(par, text, sz, bold, color, hl):
        p2 = 0
        for cm in CODE_RE.finditer(text):
            if text[p2:cm.start()]:
                style_run(par.add_run(text[p2:cm.start()]), sz, bold, color, hl)
            style_run(par.add_run(cm.group(1)), sz, bold, color, hl, code=True)
            p2 = cm.end()
        if text[p2:]:
            style_run(par.add_run(text[p2:]), sz, bold, color, hl)

    def para(text, sz=10, bold=False, indent=None, color=None,
             before=2, after=2, page_break=False, rule=False):
        p = doc.add_paragraph()
        pf = p.paragraph_format
        pf.space_before, pf.space_after = Pt(before), Pt(after)
        if indent:
            pf.left_indent = Pt(indent)
        if page_break:
            from docx.enum.text import WD_BREAK
            style_run(p.add_run(), sz)
            p.runs[0].add_break(WD_BREAK.PAGE)
        if rule:                                    # 제목 밑줄
            pbdr = docx.oxml.OxmlElement("w:pBdr")
            bot = docx.oxml.OxmlElement("w:bottom")
            for k, v in (("w:val", "single"), ("w:sz", "12"), ("w:space", "1"), ("w:color", "333333")):
                bot.set(qn(k), v)
            pbdr.append(bot)
            p._p.get_or_add_pPr().append(pbdr)
        inline(p, text, sz, bold, color)
        return p

    def set_cell_borders(tbl):
        tblPr = tbl._tbl.tblPr
        borders = docx.oxml.OxmlElement("w:tblBorders")
        for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
            e = docx.oxml.OxmlElement(f"w:{edge}")
            for k, v in (("w:val", "single"), ("w:sz", "4"), ("w:color", "8C8C8C")):
                e.set(qn(k), v)
            borders.append(e)
        tblPr.append(borders)

    def table(rows):
        ncol = max(len(r) for r in rows)
        tbl = doc.add_table(rows=len(rows), cols=ncol)
        set_cell_borders(tbl)                       # 양식에 'Table Grid' 스타일이 없어도 안전
        tbl.autofit = True
        for ri, row in enumerate(rows):
            for ci in range(ncol):
                cell = tbl.cell(ri, ci)
                txt = row[ci] if ci < len(row) else ""
                p = cell.paragraphs[0]
                p.paragraph_format.space_before = p.paragraph_format.space_after = Pt(1)
                inline(p, txt, 9, bold=(ri == 0))
                if ri == 0:                         # 머리행 음영
                    shd = docx.oxml.OxmlElement("w:shd")
                    for k, v in (("w:val", "clear"), ("w:fill", "E8EAED")):
                        shd.set(qn(k), v)
                    cell._tc.get_or_add_tcPr().append(shd)
        doc.add_paragraph().paragraph_format.space_after = Pt(3)   # 표 뒤 분리 문단

    # ── md 파싱(제품 md 부분집합) ────────────────────────────────
    lines = md_p.read_text(encoding="utf-8").split("\n")
    cells = lambda ln: [c.strip() for c in ln.strip().strip("|").split("|")]  # noqa: E731
    i, n, first_h1 = 0, len(lines), True
    while i < n:
        ln = lines[i]
        if not ln.strip() or (ln.startswith("---") and set(ln.strip()) == {"-"}):
            i += 1
            continue
        if ln.startswith("#"):
            lv = len(ln) - len(ln.lstrip("#"))
            txt = ln[lv:].strip()
            if lv == 1:
                para(txt, sz=15, bold=True, before=10, after=6,
                     page_break=not first_h1, rule=True)
                first_h1 = False
            elif lv == 2:
                para(txt, sz=12, bold=True, before=10, after=4, rule=True)
            else:
                para(txt, sz=11, bold=True, before=8, after=3)
            i += 1
            continue
        if ln.startswith("|") and i + 1 < n and re.match(r"^\|[\s:|-]+\|$", lines[i + 1]):
            rows = [cells(ln)]
            i += 2
            while i < n and lines[i].startswith("|"):
                rows.append(cells(lines[i]))
                i += 1
            table(rows)
            continue
        if ln.startswith(">"):
            while i < n and lines[i].startswith(">"):
                t = lines[i].lstrip(">").strip()
                if t:
                    para(t, sz=9, color="595959", indent=12, before=0, after=0)
                i += 1
            continue
        m = re.match(r"^(\s*)- (.*)$", ln)
        if m:
            depth = len(m.group(1)) // 2
            item = [m.group(2).strip()]
            i += 1
            while i < n and lines[i].strip() and not re.match(r"^(\s*)- ", lines[i]) \
                    and not lines[i].startswith(("#", "|", ">", "**◦")) \
                    and lines[i][:1] in (" ", "\t"):
                item.append(lines[i].strip())
                i += 1
            para(("· " if depth == 0 else "‐ ") + " ".join(item),
                 sz=10, indent=12 + depth * 12, before=1, after=1)
            continue
        buf = [ln.strip()]
        i += 1
        while i < n and lines[i].strip() and not lines[i].startswith(("#", "|", ">", "-", " ")):
            buf.append(lines[i].strip())
            i += 1
        para(" ".join(buf), sz=10, before=3, after=3)

    doc.save(str(out))

    # ── ★규칙 11 — 저장한 것을 다시 열어 확인 ────────────────────
    #   ★검증 착시 주의(2026-09-03 실제 겪음):
    #   · paragraph.text 는 하이퍼링크 안 글자를 세지 않아 본문이 적어 보인다
    #   · rels 는 같은 URL 을 1개로 중복 제거해 "링크가 빠진" 것처럼 보인다
    #   → 본문·링크 수는 XML 에서 직접 세고, 링크는 rId 정합성으로 검증한다.
    import zipfile as _zf
    chk = docx.Document(str(out))                    # 라이브러리 재개봉 자체가 1차 검증
    n_t = len(chk.tables)
    with _zf.ZipFile(out) as z:
        xml = z.read("word/document.xml").decode("utf-8")
        rels = z.read("word/_rels/document.xml.rels").decode("utf-8")
    full_text = "".join(m.group(1) for m in re.finditer(r"<w:t[^>]*>(.*?)</w:t>", xml, re.S))
    used_ids = re.findall(r'<w:hyperlink[^>]*r:id="([^"]+)"', xml)
    rel_ids = set(re.findall(r'Id="([^"]+)"[^>]*relationships/hyperlink', rels))
    dangling = [i for i in used_ids if i not in rel_ids]
    n_p = xml.count("<w:p>") + xml.count("<w:p ")
    if out.stat().st_size == 0 or not full_text.strip():
        print("❌ 산출물이 비었다")
        return 1
    if dangling:
        print(f"❌ rels 에 없는 하이퍼링크 rId {len(dangling)}개 — 워드에서 깨진 링크가 된다")
        return 1
    print(f"✅ {out} — {out.stat().st_size / 1024:.0f}KB · 문단 {n_p}(표 포함) · 표 {n_t} · "
          f"하이퍼링크 {len(used_ids)}개(고유 URL {len(rel_ids)}) · 본문 {len(full_text):,}자 "
          f"— 재개봉·rId 정합 검증 통과")
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
